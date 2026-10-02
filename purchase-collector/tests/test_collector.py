import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from purchase_collector import cli, collect, export, links, parser

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def load(name):
    return parser.load(FIXTURES / name)


class ParserTest(unittest.TestCase):
    def test_pc_format_with_multiline(self):
        msgs = load("chat_pc.txt")
        self.assertEqual(len(msgs), 10)
        self.assertEqual(msgs[0].sender, "김조교")
        self.assertEqual(msgs[0].time, datetime(2026, 9, 30, 10, 2))
        water = msgs[3]
        self.assertEqual(water.sender, "박서준")
        self.assertIn("smartstore", water.text)  # 다음 줄의 링크가 같은 메시지로 붙음
        self.assertEqual(msgs[-1].time, datetime(2026, 10, 1, 16, 0))

    def test_android_skips_system_lines(self):
        msgs = load("chat_android.txt")
        self.assertEqual([m.sender for m in msgs], ["정민호", "정민호", "한지우"])
        self.assertIn("gmarket", msgs[-1].text)

    def test_ios_and_mac(self):
        ios = load("chat_ios.txt")
        self.assertEqual(ios[0].time, datetime(2026, 10, 2, 9, 1))
        mac = load("chat_mac.csv")
        self.assertEqual([m.sender for m in mac], ["오세린", "오세린"])

    def test_noon_and_midnight(self):
        msgs = parser.parse_txt(
            "2026년 10월 1일 오후 12:05, A : x\n2026년 10월 1일 오전 12:05, B : y\n"
        )
        self.assertEqual([m.time.hour for m in msgs], [12, 0])

    def test_cp949_file(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d, "chat.txt")
            p.write_bytes((FIXTURES / "chat_ios.txt").read_text(encoding="utf-8").encode("cp949"))
            self.assertEqual(len(parser.load(p)), 1)


class LinksTest(unittest.TestCase):
    def test_same_product_normalizes_equal(self):
        a = "https://www.coupang.com/vp/products/1?itemId=2&vendorItemId=3&q=cup&rank=3"
        b = "https://m.coupang.com/vp/products/1?vendorItemId=3&itemId=2&utm_source=share#review"
        self.assertEqual(links.normalize(a), links.normalize(b))
        # 옵션(itemId)이 다르면 다른 상품
        c = "https://www.coupang.com/vp/products/1?itemId=9&vendorItemId=3"
        self.assertNotEqual(links.normalize(a), links.normalize(c))

    def test_shop_and_quantity(self):
        self.assertEqual(links.shop_of("https://smartstore.naver.com/x/products/1"), "네이버 스마트스토어")
        self.assertIsNone(links.shop_of("https://forms.gle/abc"))
        self.assertIsNone(links.shop_of("https://notcoupang.com/x"))
        self.assertEqual(links.quantity("종이컵 2박스요")[0], 2)
        self.assertEqual(links.quantity("이걸로 부탁드려요"), (None, ""))

    def test_url_trailing_punctuation(self):
        self.assertEqual(links.extract_urls("여기요 https://a.com/x."), ["https://a.com/x"])


class CollectTest(unittest.TestCase):
    def test_pc_chat(self):
        result = collect.collect(load("chat_pc.txt"))
        reqs = {r.shop: r for r in result.requests}
        self.assertEqual(len(result.requests), 3)

        cup = reqs["쿠팡"]
        self.assertEqual([a.sender for a in cup.asks], ["이하늘", "최유진"])  # 이하늘 재게시는 1번만
        self.assertEqual(cup.total_qty, 3)
        self.assertIn("종이컵 2박스요", cup.memo)  # 링크 뒤에 덧붙인 말까지 메모로

        self.assertEqual(reqs["11번가"].total_qty, 1)
        self.assertIsNone(reqs["네이버 스마트스토어"].total_qty)
        self.assertEqual([u for _, u in result.other_links], ["https://forms.gle/SampleForm123"])

    def test_period_filter(self):
        result = collect.collect(load("chat_pc.txt"), since=datetime(2026, 10, 1))
        self.assertEqual([r.shop for r in result.requests], ["11번가", "쿠팡"])

    def test_all_links_option(self):
        result = collect.collect(load("chat_pc.txt"), all_links=True)
        self.assertIn("기타", [r.shop for r in result.requests])
        self.assertEqual(result.other_links, [])

    def test_follow_up_from_other_person_is_not_merged(self):
        msgs = parser.parse_txt(
            "2026년 10월 1일 오후 3:00, A : https://www.coupang.com/vp/products/1\n"
            "2026년 10월 1일 오후 3:00, B : 저는 5개요\n"
        )
        req = collect.collect(msgs).requests[0]
        self.assertIsNone(req.total_qty)


class ExportTest(unittest.TestCase):
    def test_xlsx_and_cli(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as d:
            out = Path(d, "out.xlsx")
            self.assertEqual(cli.main([str(FIXTURES / "chat_pc.txt"), "-o", str(out), "--mask-names"]), 0)
            wb = load_workbook(out)
            ws = wb["구매요청"]
            self.assertEqual([c.value for c in ws[1]][:3], ["번호", "요청일시", "요청자"])
            self.assertEqual(ws["C2"].value, "이*늘")
            self.assertTrue(ws["H2"].value.startswith("=IF("))
            self.assertEqual(ws["A5"].value, "합계")
            self.assertIn("기타 링크", wb.sheetnames)

    def test_csv(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d, "out.csv")
            export.to_csv(collect.collect(load("chat_pc.txt")), out)
            text = out.read_text(encoding="utf-8-sig")
            self.assertIn("같은 상품 2명 요청: 이하늘(2개), 최유진(1개)", text)

    def test_mask(self):
        self.assertEqual([export.mask(n) for n in ["김", "이준", "김민수", "남궁민수"]], ["김", "이*", "김*수", "남**수"])


if __name__ == "__main__":
    unittest.main()
