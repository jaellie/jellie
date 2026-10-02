"""엑셀 구매요청 집계표 만들기. 단가만 입력하면 금액과 합계가 자동 계산된다."""

from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path

from .collect import Result

HEADERS = ["번호", "요청일시", "요청자", "쇼핑몰", "요청 내용", "수량", "단가(입력)", "금액", "링크", "처리상태", "비고"]
STATUS_OPTIONS = "대기,구매완료,보류,제외"


def mask(name: str) -> str:
    if len(name) <= 1:
        return name
    return name[0] + "*" if len(name) == 2 else name[0] + "*" * (len(name) - 2) + name[-1]


def rows(result: Result, mask_names: bool = False) -> list[list]:
    show = mask if mask_names else (lambda n: n)
    out = []
    for no, req in enumerate(result.requests, start=1):
        note = ""
        if len(req.asks) > 1:
            parts = [f"{show(a.sender)}({a.qty}개)" if a.qty else show(a.sender) for a in req.asks]
            note = f"같은 상품 {len(req.asks)}명 요청: " + ", ".join(parts)
        out.append([
            no,
            req.first.time.strftime("%Y-%m-%d %H:%M"),
            show(req.first.sender),
            req.shop,
            req.memo,
            req.total_qty,
            None,
            None,
            req.url,
            "대기",
            note,
        ])
    return out


def to_csv(result: Result, path: Path, mask_names: bool = False) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerows(rows(result, mask_names))


def to_xlsx(result: Result, path: Path, mask_names: bool = False, title: str = "") -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "구매요청"
    head_fill = PatternFill("solid", fgColor="DDE5F0")
    input_fill = PatternFill("solid", fgColor="FFF7D6")
    thin = Side(style="thin", color="BBBBBB")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)

    ws.append(HEADERS)
    for c in ws[1]:
        c.font, c.fill, c.border = Font(bold=True), head_fill, border
        c.alignment = Alignment(horizontal="center", vertical="center")

    data = rows(result, mask_names)
    for r in data:
        ws.append(r)
    last = len(data) + 1
    for i in range(2, last + 1):
        ws[f"H{i}"] = f'=IF(AND(ISNUMBER(F{i}),ISNUMBER(G{i})),F{i}*G{i},"")'
        ws[f"G{i}"].fill = input_fill
        ws[f"G{i}"].number_format = ws[f"H{i}"].number_format = "#,##0"
        link = ws[f"I{i}"]
        link.hyperlink, link.style = link.value, "Hyperlink"
        for c in ws[i]:
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=c.column_letter in "EK")

    total = last + 1
    ws[f"A{total}"] = "합계"
    ws[f"H{total}"] = f"=SUM(H2:H{last})"
    ws[f"H{total}"].number_format = "#,##0"
    for c in ws[total]:
        c.font, c.fill = Font(bold=True), head_fill

    status = DataValidation(type="list", formula1=f'"{STATUS_OPTIONS}"', allow_blank=True)
    ws.add_data_validation(status)
    if last >= 2:
        status.add(f"J2:J{last}")

    for col, width in zip("ABCDEFGHIJK", [6, 17, 10, 16, 40, 7, 11, 11, 40, 10, 30]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{last}"

    # 요약 시트
    summary = wb.create_sheet("요약")
    summary.append([title or "구매 요청 집계"])
    summary["A1"].font = Font(bold=True, size=13)
    summary.append(["읽은 메시지", result.scanned])
    summary.append(["구매 요청(상품 기준)", len(result.requests)])
    summary.append(["중복 요청이 합쳐진 상품", sum(1 for r in result.requests if len(r.asks) > 1)])
    summary.append(["예상 총액", "='구매요청'!H" + str(total)])
    summary["B5"].number_format = "#,##0"
    summary.append([])
    summary.append(["요청자", "건수"])
    show = mask if mask_names else (lambda n: n)
    for name, n in Counter(show(a.sender) for r in result.requests for a in r.asks).most_common():
        summary.append([name, n])
    summary.append([])
    summary.append(["쇼핑몰", "건수"])
    for shop, n in Counter(r.shop for r in result.requests).most_common():
        summary.append([shop, n])
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 14

    # 쇼핑몰이 아닌 링크 (공지, 설문 등) — 놓친 요청이 없는지 사람이 훑어볼 수 있게
    if result.other_links:
        other = wb.create_sheet("기타 링크")
        other.append(["일시", "보낸 사람", "링크", "메시지"])
        for msg, url in result.other_links:
            other.append([msg.time.strftime("%Y-%m-%d %H:%M"), show(msg.sender), url, msg.text[:200]])
        for col, width in zip("ABCD", [17, 10, 45, 50]):
            other.column_dimensions[col].width = width

    wb.save(path)
